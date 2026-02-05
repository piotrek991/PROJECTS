import time

import polars as pl
import asyncio
import httpx
from httpx import HTTPStatusError
from lxml import etree
import requests
import json
import re
from datetime import datetime
import os
from pathlib import Path
from concurrent.futures import ProcessPoolExecutor, as_completed, ThreadPoolExecutor
import pprint
import openpyxl
from openpyxl.styles import PatternFill, Alignment
import xlsxwriter
from openpyxl import load_workbook
from typing import Dict, Any, Coroutine, List
from enum import Enum
import json
from functools import reduce
from tenacity import (
    retry,
    stop_after_attempt,
    wait_exponential,
    retry_if_exception_type,
    retry_if_result
)
from logger import setup_custom_logger
from functools import wraps
import inspect
import random
import uuid, base64, string
import requests
import http.client
import urllib.parse
from requests.cookies import RequestsCookieJar
from seleniumwire import webdriver
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.common.exceptions import TimeoutException, WebDriverException

_MIN_BASE_WIDTH_CELL=10
_MAX_BASE_HEIGHT_CELL=150
countries_data = dict()

global_logger = setup_custom_logger('DATA PLAYERS')
global_sem = asyncio.Semaphore(2)

def _gen_new_cookie():
    def generate_random_string(length):
        return ''.join(random.choices(string.ascii_letters + string.digits, k=length))

    def generate_random_base64(length):
        random_bytes = ''.join(random.choices(string.ascii_letters + string.digits, k=length)).encode()
        return base64.b64encode(random_bytes).decode()

    # Simplified valid IAB TCF v2 consent string
    consent_string = "CPi0AAAPi0AAAAENAACAAAAAAAAAAAAAAAAAAAAA.YAAAAAAAAAA"

    # Generate a random _sp_v1_ss value mimicking the provided format
    session_base64 = "H4sIAAAAAAAAA" + generate_random_base64(50)
    sp_v1_ss = f"1:{session_base64}"
    url_encoded_sp_v1_ss = urllib.parse.quote(sp_v1_ss, safe=":")  # URL-encode, keeping ":" safe

    cookies_dict = {
        "_sp_su": "false",
        "_sp_v1_p": '476',#str(random.randint(1, 1000)),
        "_sp_v1_data": '745500',#str(random.randint(100000, 999999)),
        "_sp_v1_ss": '1:H4sIAAAAAAAAAItWqo5RKimOUbKKhjHySnNydGKUUpHYJWCJ6traWFwSSjqYBmFn5IEYBrhNoqKEUiwAqTzxSd8AAAA%3D',#url_encoded_sp_v1_ss,
        "euconsent-v2": consent_string,
        "consentUUID": '711d0124-2b29-46f2-8344-7c5e3d713538_44_45_47',#f"{uuid.uuid4()}_44_45_47",
        "cuukie": 'a0ZOVUdOeTRtdkZJbndafjg0aTNCfnRGQ1RSc1lNWEoEpg2UlBaZ0oN6FoMLcVuaEaSXxLHWhrPyY4FNZOzMjA%3D%3D',#generate_random_base64(50)
    }

    # Create RequestsCookieJar
    cookiejar = RequestsCookieJar()
    for name, value in cookies_dict.items():
        cookiejar.set(name, value, domain="www.transfermarkt.com")

    # Log cookie string for debugging
    cookie_string = "; ".join(f"{key}={value}" for key, value in cookies_dict.items())
    global_logger.debug(f"Generated cookie string: {cookie_string}")

    return cookiejar

def _choose_random_ua() -> str:
    ua_l = [
        'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/139.0.7113.106 Safari/537.36'
        , 'Mozilla/5.0 (Windows NT 10.0; Win64; x64; rv:129.0) Gecko/20100101 Firefox/129.0'
        , 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/140.0.7240.88 Safari/537.36'
        , 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/138.0.6998.81 Safari/537.36 Edg/138.0.2485.51'
        , 'Mozilla/5.0 (Windows NT 10.0; Win64; x64; rv:128.0) Gecko/20100101 Firefox/128.0'
        , 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/141.0.7321.62 Safari/537.36'
        , 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/137.0.6952.40 Safari/537.36 OPR/123.0.0.0'
        , 'Mozilla/5.0 (Windows NT 10.0; Win64; x64; rv:130.0) Gecko/20100101 Firefox/130.0'
        , 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/140.0.7245.19 Safari/537.36 Edg/140.0.2600.32'
        , 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/136.0.6880.72 Safari/537.36'
        , 'Mozilla/5.0 (Windows NT 10.0; Win64; x64; rv:127.0) Gecko/20100101 Firefox/127.0'
        , 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/135.0.6843.98 Safari/537.36 Vivaldi/6.9.999.41'
        , 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/139.0.7100.55 Safari/537.36 OPR/124.0.0.0'
        , 'Mozilla/5.0 (Windows NT 10.0; Win64; x64; rv:131.0) Gecko/20100101 Firefox/131.0'
        , 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/134.0.6765.43 Safari/537.36'
        , 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/141.0.7318.10 Safari/537.36 Edg/141.0.2701.9'
        , 'Mozilla/5.0 (Windows NT 10.0; Win64; x64; rv:126.0) Gecko/20100101 Firefox/126.0'
        , 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/133.0.6712.12 Safari/537.36'
        , 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/138.0.7002.5 Safari/537.36 Edg/138.0.2490.5'
        , 'Mozilla/5.0 (Windows NT 10.0; Win64; x64; rv:125.0) Gecko/20100101 Firefox/125.0'
    ]
    return ua_l[random.randint(0, len(ua_l)-1)]

def _url_logger_wrapper(func):
    sig = inspect.signature(func)
    @wraps(func)
    def wrapper(*args, **kwargs):
        bound = sig.bind(*args, **kwargs)
        url = bound.arguments.get('url')
        country = bound.arguments.get('country_to_log')
        global_logger.debug(f"PROCESSING {url} FOR COUNTRY {country} IN FUNCTION {func.__name__}")
        return func(*args, **kwargs)
    return wrapper


class JoinHow(str, Enum):
    INNER = "inner"
    LEFT  = "left"
    OUTER = "outer"
    SEMI  = "semi"
    ANTI  = "anti"
    CROSS = "cross"

class OperationalError(Exception):
    def __init__(self, message: str, payload: Dict[str, Any] | None = None) -> None:
        self.message = message
        self.payload = payload or {}
        super().__init__(message)

def get_f_list(d_path: Path) -> list | None:
    if d_path.is_dir():
        f_list = list(d_path.glob('*.xlsx'))
        return f_list
    else:
        return None

def _process_performance_data(data_j_raw: Any, outer_d: Dict[str,Any] = None) -> None | List[Dict[str,Any]]:
    try:
        perf_list = list()

        goalkeeper_b = data_j_raw.get('goalkeeper')
        for perf in data_j_raw.get('performances'):
            perf_list.append({
                'scoring': {'assists': perf.get('assists')
                    , 'gamesPlayed': perf.get('gamesPlayed')
                    , 'goalsScored': perf.get('goalsScored')
                    , 'cleanSheets': perf.get('cleanSheets')
                            }
                , 'entity': perf.get('entity', {}).get('name')
                , 'goalkeeper': goalkeeper_b
                , **(outer_d if outer_d else {})
            })
        return perf_list
    except Exception as e:
        print(f"unexpected error occurred when processing performance data: {e}")
        return None

@retry(
    stop=stop_after_attempt(3),
    wait=wait_exponential(multiplier=1, min=1, max=10),
    retry=(
        retry_if_exception_type(httpx.RequestError)
    ),
    reraise=True
)
@_url_logger_wrapper
async def _process_performance(client_i:httpx.AsyncClient, url: str, outer_d: Dict[str,str], country_to_log:str) -> None | dict[str, bool | str]:
    try:
        data_r = await client_i.get(url, follow_redirects=True)
        data_r.raise_for_status()
        data_j_processed = await asyncio.to_thread(_process_performance_data, data_r.json(), outer_d)
        return data_j_processed
    except httpx.HTTPStatusError as e:
        print(f"An unexpected HTTPStatusError error occurred for {url}: {e}")
        raise HTTPStatusError(str(e)) from e
    except httpx.RequestError as e:
        print(f"An unexpected RequestError error occurred for {url}: {e}")
        raise httpx.RequestError(str(e))
    except Exception as e:
        print(f"An unexpected error occurred for {url}: {e}")
        raise e

@retry(
    stop=stop_after_attempt(3),
    wait=wait_exponential(multiplier=1, min=1, max=10),
    retry=(
        retry_if_exception_type(httpx.RequestError)
    ),
    reraise=True
)
async def _simple_fetch(client_i:httpx.AsyncClient, url: str) -> Dict[str, Any]:
    try:
        data_r = await client_i.get(url, follow_redirects=True)
        data_r.raise_for_status()
        data_r_json = data_r.json()

        return data_r_json
    except httpx.HTTPStatusError as e:
        print(f"An unexpected HTTPStatusError error occurred for {url}: {e}")
        raise
    except httpx.RequestError as e:
        print(f"An unexpected RequestError error occurred for {url}: {e}")
        raise
    except Exception as e:
        print(f"An unexpected error occurred for {url}: {e}")
        raise

@retry(
    stop=stop_after_attempt(3),
    wait=wait_exponential(multiplier=1, min=1, max=10),
    retry=(
        retry_if_exception_type(httpx.RequestError)
    ),
    reraise=True
)
@_url_logger_wrapper
async def _process_national(client_i:httpx.AsyncClient, url: str, outer_d: Dict[str,str], country_to_log:str) -> None | Dict[str, Any]:
    club_n_url = f"https://tmapi-alpha.transfermarkt.technology/clubs?ids[]={{team_id}}"
    try:
        data_r = await client_i.get(url, follow_redirects=True)
        data_r.raise_for_status()
        data_r_json = data_r.json()

        club_n_tasks = list()
        national_hist_l = list()
        for perf_hist in data_r_json.get('data', {}).get('history'):
            club_n_tasks.append(
                asyncio.create_task(_simple_fetch(client_i, club_n_url.format(team_id=perf_hist.get('clubId'))))
            )
            national_hist_l.append({
                'gamesPlayed': perf_hist.get('gamesPlayed'),
                'goalScored': perf_hist.get('goalsScored')
            })

        if national_hist_l:
            club_n_names = await asyncio.gather(*club_n_tasks)
            for num, club_n_data in enumerate(club_n_names):
                club_n_inner = club_n_data.get('data')[0].get('name')
                national_hist_l[num].update({
                    'national_name': club_n_inner
                })
        return {**(outer_d if outer_d else {}), 'national_career_details': national_hist_l}
    except httpx.HTTPStatusError as e:
        print(f"An unexpected HTTPStatusError error occurred for {url}: {e}")
        raise
    except httpx.RequestError as e:
        print(f"An unexpected RequestError error occurred for {url}: {e}")
        raise
    except Exception as e:
        print(f"An unexpected error occurred for {url}: {e}")
        raise

@retry(
    stop=stop_after_attempt(3),
    wait=wait_exponential(multiplier=1, min=1, max=10),
    retry=(
        retry_if_exception_type(httpx.RequestError)
    ),
    reraise=True
)
@_url_logger_wrapper
async def _process_trophies_history(client_i:httpx.AsyncClient, url: str, outer_d: Dict[str,str], country_to_log:str) -> None | Dict[str, Any]:
    try:
        data_r = await client_i.get(url, follow_redirects=True)
        data_r.raise_for_status()
        data_r_html = etree.HTML(data_r.text)

        trophies_d_raw = {}
        trophies_d_processed = {}

        last_referenced = str()
        for row in data_r_html.xpath('//tbody[tr[contains(@class, "bg_Sturm")]]/tr'):
            if row.get('class') == 'bg_Sturm':
                raw_t = row.xpath('./td/text()')[0]
                trophies_d_processed.update({
                    'trophies_name': last_referenced
                    , 'years_achieved': trophies_d_raw.get(last_referenced)
                })
                last_referenced = raw_t
                trophies_d_raw.update({raw_t: []})
            else:
                t_year,t_team = (str(), str())
                for row_inner in row.xpath('./td'):
                    if check_rel_list(row_inner.xpath('./text()'),0):
                        t_year = row_inner.xpath('./text()')[0].strip()
                    elif check_rel_list(row_inner.xpath('./a/text()'),0):
                        t_team = row_inner.xpath('./a/text()')[0].strip()
                trophies_d_raw[last_referenced].append(
                    {'year': t_year, 'team': t_team}
                )
        return {**outer_d, 'trophies': trophies_d_processed}
    except httpx.HTTPStatusError as e:
        print(f"An unexpected HTTPStatusError error occurred for {url}: {e}")
        raise
    except httpx.RequestError as e:
        print(f"An unexpected RequestError error occurred for {url}: {e}")
        raise
    except Exception as e:
        print(f"An unexpected error occurred for {url}: {e}")
        raise

def check_rel_list(data_l : list, ind: int):
    try:
        data_l_ind = data_l[ind].strip()
        return data_l_ind if data_l_ind else None
    except IndexError as ind_e:
        return None

@retry(
    stop=stop_after_attempt(3),
    wait=wait_exponential(multiplier=1, min=1, max=10),
    retry=(
        retry_if_exception_type(httpx.RequestError)
    ),
    reraise=True
)
@_url_logger_wrapper
async def _process_player_general(client_i:httpx.AsyncClient, url: str, outer_d: Dict[str,str], country_to_log: str) -> None | Dict[str, Any]:
    club_n_url = f"https://tmapi-alpha.transfermarkt.technology/clubs?ids[]={{team_id}}"
    try:
        data_r = await _simple_fetch(client_i, url)
        data_r_json = data_r.get('data')[0]
        player_data = {
            'age': data_r_json.get('lifeDates', {}).get('age')
            , 'dateOfBirth': data_r_json.get('lifeDates', {}).get('dateOfBirth')
            , 'placeOfBirth': data_r_json.get('birthPlaceDetails', {}).get('placeOfBirth')
            , 'countryOfBirth': countries_data.get(data_r_json.get('birthPlaceDetails', {}).get('countryOfBirthId'))
            , 'preferredFoot': data_r_json.get('attributes', {}).get('preferredFoot', {}).get('name')
            , 'position': data_r_json.get('attributes', {}).get('position')
            , 'height': data_r_json.get('attributes', {}).get('height')
            , 'lastContractRenewal': data_r_json.get('attributes', {}).get('lastContractRenewal')
            , 'formerClubs': data_r_json.get('attributes', {}).get('formerClubsNote')
            , 'contractUntil': data_r_json.get('attributes', {}).get('contractUntil')
        }
        clubs_assignments = list()
        for team_assigment in data_r_json.get('clubAssignments'):
            data_club = await _simple_fetch(client_i, club_n_url.format(team_id=team_assigment.get('clubId')))
            club_name = data_club.get('data')[0].get('name')
            club_primary_league = data_club.get('data')[0].get('baseDetails').get('primaryCompetitionId')
            clubs_assignments.append({
                'clubId': team_assigment.get('clubId'),
                'clubName': club_name,
                'shirtNumber': team_assigment.get('shirtNumber'),
                'type': team_assigment.get('type'),
                'debut': team_assigment.get('debut'),
                'primary_competition': club_primary_league
            })
        player_data.update({
            'clubsAssignments': clubs_assignments
            , **(outer_d if outer_d else {})
        })
        return player_data
    except httpx.HTTPStatusError as e:
        print(f"An unexpected HTTPStatusError error occurred for {url}: {e}")
        raise
    except httpx.RequestError as e:
        print(f"An unexpected RequestError error occurred for {url}: {e}")
        raise
    except Exception as e:
        print(f"An unexpected error occurred for {url}: {e}")
        raise

async def get_process_t_data(client_i: httpx.AsyncClient, data_l: dict, async_sem_inner: asyncio.Semaphore, country_to_log:str) -> dict:
    try:
        url_tasks = list()
        player_name_link = re.search(r'^/(.*?)/', data_l.get('player_link'))
        if player_name_link:
            player_name_link = player_name_link.group(1)
        else:
            raise OperationalError('No player name in player name link defined', {'data': data_l})

        performance_per_comp = f"https://www.transfermarkt.com/ceapi/player/{data_l.get('player_id')}/performancepercompetition"
        performance_per_club = f"https://www.transfermarkt.com/ceapi/player/{data_l.get('player_id')}/performanceperclub"
        national_career = f"https://tmapi-alpha.transfermarkt.technology/player/{data_l.get('player_id')}/national-career-history"
        trophies_history = f"https://www.transfermarkt.com/{player_name_link}/erfolge/spieler/{data_l.get('player_id')}"
        player_data = f"https://tmapi-alpha.transfermarkt.technology/players?ids[]={data_l.get('player_id')}"
        async with async_sem_inner:
            url_tasks.extend([
                asyncio.create_task(_process_performance(client_i, performance_per_comp,
                                                         {'player_id': data_l.get('player_id'),
                                                          'player_name': data_l.get('player_name'),
                                                          'ent_name': 'competition'}, country_to_log)),
                asyncio.create_task(_process_performance(client_i, performance_per_club,
                                                         {'player_id': data_l.get('player_id'),
                                                          'player_name': data_l.get('player_name'),
                                                          'ent_name': 'club'}, country_to_log)),
                asyncio.create_task(_process_national(client_i, national_career, {'player_id': data_l.get('player_id'),
                                                          'player_name': data_l.get('player_name'),
                                                          'ent_name': 'national_career'}, country_to_log)),
                asyncio.create_task(_process_trophies_history(client_i, trophies_history, {'player_id': data_l.get('player_id'),
                                                          'player_name': data_l.get('player_name'),
                                                          'ent_name': 'trophies_history'}, country_to_log)),
                asyncio.create_task(_process_player_general(client_i, player_data, {'player_id': data_l.get('player_id'),
                                                                           'player_name': data_l.get('player_name'),
                                                                           'ent_name': 'player_details'}, country_to_log)),
            ])
            data_combined_result = await asyncio.gather(*url_tasks)
            return  { 'process_result': True
                , 'data_comp':data_combined_result[0]
                , 'data_club':data_combined_result[1]
                , 'data_national':data_combined_result[2]
                , 'data_trophies':data_combined_result[3]
                , 'data_player':data_combined_result[4]
                     }
    except httpx.HTTPStatusError as e:
        print(f"HTTP Status Error for {data_l.get('player_id')}: {e.response.status_code} - {e.response.text.strip()[:100]}...")
        return {'process_result': False, 'fetched_data': e}
    except httpx.RequestError as e:
        print(f"Request Error for {data_l.get('player_id')}: {e} and  {data_combined_result}")
        return {'process_result': False, 'fetched_data': e}
    except Exception as e:
        print(f"An unexpected error occurred for {data_l}: {e}")
        return {'process_result': False, 'fetched_data': e}


@retry(
    stop=stop_after_attempt(3),
    wait=wait_exponential(multiplier=1, min=1, max=10),
    retry=(
        retry_if_exception_type(httpx.RequestError) |
        retry_if_exception_type(httpx.HTTPStatusError)
    ),
    reraise=True
)
@_url_logger_wrapper
async def _process_leagues(client_i: httpx.AsyncClient, url:str, outer_d: Dict[str,str], sem: asyncio.Semaphore) ->  None | Dict[str, Any]:
    async with sem:
        try:
            data_response = await _simple_fetch(client_i, url)
            data_j = data_response.get('data')[0]
            f_d = {'league_level': data_j.get('typeId'), 'league_name': data_j.get('name')}

            return {**f_d, **outer_d}
        except httpx.HTTPStatusError as e:
            print(f"An unexpected HTTPStatusError error occurred for {url}: {e}")
            raise
        except httpx.RequestError as e:
            print(f"An unexpected RequestError error occurred for {url}: {e}")
            raise
        except Exception as e:
            print(f"An unexpected error occurred for {url}: {e}")
            raise

@retry(
    stop=stop_after_attempt(3),
    wait=wait_exponential(multiplier=2, min=5, max=20),
    retry=retry_if_exception_type((TimeoutException, WebDriverException)),
    reraise=True
)
@_url_logger_wrapper
async def _process_leagues_selenium(url: str, outer_d: Dict[str, str], sem: asyncio.Semaphore) -> None | Dict[str, Any]:
    async with sem:
        global_logger.debug(f"Semaphore acquired for {url}, active count: {sem._value}")
        try:
            def sync_process():
                chrome_options = Options()
                chrome_options.add_argument(f"user-agent=Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/139.0.0.0 Safari/537.36")
                chrome_options.add_argument("--headless")
                chrome_options.add_argument("--disable-gpu")
                chrome_options.add_argument("--no-sandbox")
                chrome_options.add_argument("--disable-dev-shm-usage")
                options = {
                    'accept': 'text/html,application/xhtml+xml,application/xml;q=0.9,image/avif,image/webp,image/apng,*/*;q=0.8,application/signed-exchange;v=b3;q=0.7',
                    'accept-language': 'pl-PL,pl;q=0.9,en-US;q=0.8,en;q=0.7',
                    'cache-control': 'no-cache',
                    'pragma': 'no-cache',
                    'priority': 'u=0, i',
                    'referer': url,
                    'sec-ch-ua': '"Not;A=Brand";v="99", "Google Chrome";v="139", "Chromium";v="139"',
                    'sec-ch-ua-mobile': '?0',
                    'sec-ch-ua-platform': '"Windows"',
                    'sec-fetch-dest': 'document',
                    'sec-fetch-mode': 'navigate',
                    'sec-fetch-site': 'same-origin',
                    'sec-fetch-user': '?1',
                    'upgrade-insecure-requests': '1'
                }

                driver = webdriver.Chrome(options=chrome_options, seleniumwire_options=options)
                try:
                    driver.get("https://www.transfermarkt.com")
                    cookiejar = _gen_new_cookie()
                    for cookie in cookiejar:
                        driver.add_cookie({"name": cookie.name, "value": cookie.value, "domain": "www.transfermarkt.com"})

                    driver.get(url)
                    WebDriverWait(driver, 10).until(
                        EC.presence_of_element_located((By.XPATH, "//div[@class='data-header__headline-container']/h1"))
                    )
                    data_r_html = etree.HTML(driver.page_source)

                    league_name = data_r_html.xpath("//div[@class='data-header__headline-container']/h1/text()")[0].strip()
                    f_d = {'league_level': None, 'league_name': league_name}
                    for item in data_r_html.xpath("//div[@class='data-header__club-info']/span[@class='data-header__label']"):
                        if check_rel_list(item.xpath('./text()'), 0):
                            item_t = item.xpath('./text()')[0].strip()
                            if item_t == 'League level:':
                                l_level = item.xpath('./span/text()')[0].strip()
                                f_d.update({'league_level': l_level})

                    #time.sleep(random.uniform(3.0, 6.0))
                    return {**f_d, **outer_d}
                finally:
                    driver.quit()

            result = await asyncio.to_thread(sync_process)
            global_logger.debug(f"Semaphore released for {url}")
            return result
        except TimeoutException as e:
            global_logger.warning(f"Timeout for {url}, waiting 20 seconds")
            time.sleep(20)
            raise
        except WebDriverException as e:
            global_logger.error(f"WebDriverException for {url}: {e}")
            raise
        except Exception as e:
            global_logger.error(f"Unexpected error for {url}: {e}")
            raise

async def main_wfl(f_file_to_process: Path, only_main_leagues: bool = False):
    async_sem = asyncio.Semaphore(3)
    page_headers = {
        #'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/139.0.0.0 Safari/537.36 Edg/139.0.0.0 Agency/98.8.7987.78',
    'accept': 'text/html,application/xhtml+xml,application/xml;q=0.9,image/avif,image/webp,image/apng,*/*;q=0.8,application/signed-exchange;v=b3;q=0.7',
    'cache-control': 'no-cache',
    'pragma': 'no-cache',
    'priority': 'u=0, i',
    'sec-ch-ua': 'Not;A=Brand";v="99", "Google Chrome";v="139", "Chromium";v="139"',
    'sec-ch-ua-mobile': '?0',
    'sec-ch-ua-platform': '"Windows"',
    'sec-fetch-dest': 'document',
    'sec-fetch-mode': 'navigate',
    'sec-fetch-site': 'same-origin',
    'sec-fetch-user': '?1',
    'user-agent':  _choose_random_ua()
    , "Accept-Language": "en-US,en;q=0.9"
    }
    try:
        country_nm = re.search(r'data_(.*?).xlsx', str(f_file_to_process.name)).group(1)
        global_logger.info(f"STARTING TO PROCESS COUNTRY {country_nm}")

        pl_df_raw = (pl.read_excel(f_file_to_process, sheet_name=country_nm))

        if only_main_leagues:
            d_leagues = pl_df_raw.select('league_id').drop_nulls().unique().to_series().to_list()
            async with httpx.AsyncClient(headers=page_headers, timeout=60) as client:
                tasks_league = list()
                for league in d_leagues:
                    tasks_league.append(_process_leagues(client, f'https://tmapi-alpha.transfermarkt.technology/competitions?ids[]={league}', {'league_id': league}, async_sem))
                data_leagues = await asyncio.gather(*tasks_league)

            leagues_df = pl.DataFrame(data_leagues)
            global_logger.debug(f"PROCESSING WITH LEAGUES DATA {leagues_df}")
            pl_df = (
                pl_df_raw.join(leagues_df, on='league_id', how='left').filter(pl.col('league_level') == 1)
                     .select('player_name','player_id', 'player_link', 'league_name', 'league_level', 'league_id').drop_nulls().unique()
            )
        else:
            pl_df = (
                pl_df_raw.select('player_name', 'player_id', 'player_link').drop_nulls().unique()
            )
        pl_df_r = pl_df.select(pl.struct(['player_id', 'player_link', 'player_name']).alias('t_data')).to_series().to_list()
        #pl_df_r = pl_df_r[:10]
        global_logger.debug(f"PLAYERS TO PROCESS {len(pl_df_r)} with {pl_df}")


        async with httpx.AsyncClient(headers=page_headers, timeout=60) as client:
            await _set_countries_data(client)
            tasks = list()
            for player_t in pl_df_r:
                tasks.append(asyncio.create_task(get_process_t_data(client, player_t, async_sem, country_nm)))
            data_f = await asyncio.gather(*tasks, return_exceptions=True)
            print(f"AFTER RESULT")
        result_d = {
            'data_comp': []
            , 'data_club': []
            , 'data_national': []
            , 'data_trophies': []
            , 'data_player': []
        }
        exceptions_l = list()
        for df_result in data_f:
            try:
                result_d['data_comp'].extend(df_result['data_comp'])
                result_d['data_club'].extend(df_result['data_club'])
                result_d['data_national'].append(df_result['data_national'])
                result_d['data_trophies'].append(df_result['data_trophies'])
                result_d['data_player'].append(df_result['data_player'])
            except Exception as e:
                global_logger.warning(f"DATA FOR ONE PLAYER SKIPPED BECAUSE OF ERROR - {e} FOR COUNTRY {country_nm}")
                exceptions_l.append(df_result)
                continue

        pl_df_comp = (pl.DataFrame(result_d.get('data_comp')).lazy()
        .unnest('scoring').with_columns(
            pl.when(~pl.col('goalkeeper')).then(pl.format('Competition Name: {} - Scoring - Games Played: {} | Assists: {} | Goals: {}', pl.col('entity'), pl.col('gamesPlayed'), pl.col('assists'), pl.col('goalsScored'))) \
                .otherwise(pl.format('Competition Name: {} - Scoring - Clean Sheets: {}',pl.col('entity'), pl.col('cleanSheets'))).alias('COMPETITION HISTORY DETAILS')
        ).select(pl.col('player_id'), pl.col('player_name'), pl.col('COMPETITION HISTORY DETAILS')).unique()
        .group_by(['player_id', 'player_name']).agg(pl.col('COMPETITION HISTORY DETAILS').str.join('\n,'))
                      )
        pl_df_club = (
            pl.DataFrame(result_d.get('data_club')).lazy()
            .unnest('scoring').with_columns(
                pl.when(~pl.col('goalkeeper')).then(
                    pl.format('Club Name: {} - Scoring - Games Played: {} | Assists: {} | Goals: {}',
                              pl.col('entity'), pl.col('gamesPlayed'), pl.col('assists'), pl.col('goalsScored')))
                    .otherwise(pl.format('Club Name: {} - Scoring - Clean Sheets: {}', pl.col('entity'),
                                         pl.col('cleanSheets'))).alias('CLUB HISTORY DETAILS')
            ).select(pl.col('player_id'), pl.col('player_name'), pl.col('CLUB HISTORY DETAILS')).unique()
            .group_by(['player_id', 'player_name']).agg(pl.col('CLUB HISTORY DETAILS').str.join('\n,'))
        )
        pl_df_national = (
            pl.DataFrame(result_d.get('data_national')).lazy()
                .explode('national_career_details').unnest('national_career_details').with_columns(
                    pl.format('National Team Name: {} - Scoring - Games Played: {} | Goals: {}', pl.col('national_name'), pl.col('gamesPlayed'), pl.col('goalScored') )
                        .fill_null('NO DATA OF PLAYING IN NATIONAL TEAM').alias('NATIONAL TEAM HISTORY DETAILS')
            ).unique().group_by(['player_id', 'player_name']).agg(pl.col('NATIONAL TEAM HISTORY DETAILS').str.join('\n,'))
        )

        pl_df_trophies = (
            pl.DataFrame(result_d.get('data_trophies')).lazy()
                .unnest('trophies').explode('years_achieved').with_columns(
                    pl.col('years_achieved').struct.with_fields(
                        pl.format('Achievement Name: {} - Season Achieved: {} | Team Name: {}', pl.col('trophies_name'), pl.field('year'), pl.field('team')).alias('ACHIEVEMENTS HISTORY DETAILS')
                    ).struct.field('ACHIEVEMENTS HISTORY DETAILS').str.replace_all('^\\s*$','NO RECORDED ACHIEVEMENTS').fill_null('NO RECORDED ACHIEVEMENTS').alias('ACHIEVEMENTS HISTORY DETAILS')
            )
        ).select('player_id', 'player_name', 'ACHIEVEMENTS HISTORY DETAILS').unique().group_by(['player_id', 'player_name']).agg(pl.col('ACHIEVEMENTS HISTORY DETAILS').str.join('\n,'))

        pl_df_player = (
            pl.DataFrame(result_d.get('data_player')).lazy()
            .explode('clubsAssignments').with_columns(
                pl.col(pl.String).fill_null('NOT SPECIFIED')
                , pl.col(pl.Float64).fill_null(0.0)
            ).with_columns(
                pl.col('clubsAssignments').struct.with_fields(
                    pl.format('Club Name: {} | Shirt Number: {} | Debut Date {} | Type: {}'
                              , pl.coalesce(pl.field('clubName'), pl.lit('NOT SPECIFIED'))
                              , pl.coalesce(pl.field('shirtNumber'), pl.lit('NOT SPECIFIED'))
                              , pl.coalesce(pl.field('debut'), pl.lit('NOT SPECIFIED'))
                              , pl.coalesce(pl.field('type'), pl.lit('NOT SPECIFIED'))
                              ).alias('CLUB ASSIGNMENTS DETAILS')
                ).struct.field('CLUB ASSIGNMENTS DETAILS')
                , pl.col('clubsAssignments').struct.field('primary_competition')
                , pl.col('placeOfBirth').str.replace_all(r"^$", "NOT SPECIFIED")
                , pl.col('dateOfBirth').cast(pl.Date, strict=False).fill_null('NOT SPECIFIED')
                , pl.col('lastContractRenewal').struct.with_fields(
                    pl.concat_str([
                        pl.field('year')
                        , pl.field('month')
                        , pl.field('day')
                    ], separator='-').str.strptime(pl.Date, '%Y-%m-%d', strict=False).fill_null('NOT SPECIFIED').alias('lContractRenewal')
                ).struct.field('lContractRenewal')
                , pl.col('contractUntil').str.strptime(pl.Datetime, '%+', strict=False).dt.date().fill_null('NOT SPECIFIED')
                , pl.col('position').struct.field('name').fill_null('NOT SPECIFIED').alias('position')
            ).unique().group_by(['player_id', 'player_name']).agg(
                pl.col('CLUB ASSIGNMENTS DETAILS').str.join('\n,')
                ,pl.col('placeOfBirth').first()
                , pl.col('dateOfBirth').first()
                , pl.col('countryOfBirth').first()
                , pl.col('lContractRenewal').first()
                , pl.col('height').first()
                , pl.col('age').first()
                , pl.col('preferredFoot').first()
                , pl.col('position').first()
                , pl.col('primary_competition').first()
            )
        )

        df_combined = (_multi_join_polars([pl_df_player, pl_df_comp, pl_df_national, pl_df_trophies, pl_df_club], on_field=['player_id', 'player_name'])
                       .rename({'player_name': 'Player Name'
                            , 'placeOfBirth': 'Place of Birth'
                            , 'dateOfBirth': 'Date of Birth'
                            , 'countryOfBirth': 'Country of Birth'
                            , 'lContractRenewal': 'Last Contract Renewal'
                            , 'height': 'Height'
                            , 'age': 'Age'
                            , 'preferredFoot': 'Preferred Foot'
                            , 'position': 'Position'
                            , 'CLUB ASSIGNMENTS DETAILS': 'Club Assignments'
                            , 'COMPETITION HISTORY DETAILS': 'Competition History Details'
                            , 'NATIONAL TEAM HISTORY DETAILS': 'National Team History Details'
                            , 'CLUB HISTORY DETAILS': 'Club History Details'
                            , 'ACHIEVEMENTS HISTORY DETAILS': 'Achievements History Details'
                            }
                           ).collect())
        if only_main_leagues:
            df_res = pl_df.select('player_id', 'league_name', 'league_level', 'league_id')
            df_combined = (
                df_combined.join(df_res, left_on=['player_id', 'primary_competition'], right_on=['player_id', 'league_id'], how='inner')
                .drop('primary_competition')
                .rename({'league_name': 'Current League Name', 'league_level': 'Current League Level'})
            )
        df_combined = df_combined.drop('player_id')

        d_name =  f'players_data_details{"_first_tier" if only_main_leagues else ""}'
        os.makedirs(os.path.join(os.getcwd(), d_name), exist_ok=True)
        f_name_base = f"players_data_{country_nm}.xlsx"
        df_combined.write_excel(
            os.path.join( d_name, f_name_base)
            , 'data'
            , dtype_formats={
                pl.Float32: "#,##0.00"
            }
        )
        wb = openpyxl.load_workbook(os.path.join( d_name, f_name_base))
        sheet = wb.active
        default_row_height = sheet.row_dimensions[1].height or 15

        col_to_center = [
            'Club Assignments'
            , 'Competition History Details'
            , 'National Team History Details'
            , 'Achievements History Details'
            , 'Club History Details'
        ]
        cells_to_center = [chr(65 + df_combined.columns.index(item)) for item in col_to_center]
        pattern = PatternFill(start_color='ADD8E6', end_color='ADD8E6', fill_type='solid')
        letter_to_num = chr(65 + df_combined.columns.index('Height'))
        for cell in sheet[1]:
            cell.fill = pattern

        sheet.freeze_panes = 'A2'
        row_height_ref = {}
        for loc, letter in enumerate([chr(65 + i) for i in range(len(df_combined.columns))], start=1):
            m_width = _MIN_BASE_WIDTH_CELL
            for r_num in range(1, sheet.max_row + 1):
                m_height = default_row_height
                if letter == letter_to_num:
                    sheet.cell(row=r_num, column=loc).number_format = '#,##0.00'
                if sh_option := sheet[f'{letter}{r_num}'].value:
                    if type(sh_option) == str:
                        sh_option = sh_option.replace('\r\n', '\n').replace('\r','\n').split(',')
                        sh_option_m_len = reduce(lambda a,b: max(a, len(b)), sh_option, len(sh_option[0]))
                        sheet.cell(row=r_num, column=loc).alignment = Alignment(
                            horizontal="center" if (letter in cells_to_center) and r_num > 1 else "left",
                            vertical="center" ,
                            wrap_text=True
                        )
                        if sh_option_m_len > m_width:
                            m_width = sh_option_m_len
                        if (len(sh_option) * default_row_height) > m_height:
                            m_height = len(sh_option) * default_row_height
                if row_height_ref.get(r_num, 0) < m_height:
                    row_height_ref[r_num] = m_height
                    sheet.row_dimensions[r_num].height = m_height
            sheet.column_dimensions[letter].width = m_width

        wb.save(os.path.join( d_name, f_name_base))
        wb.close()

        return {'file': f_file_to_process, 'processed': True, 'exceptions': exceptions_l}
    except Exception as e:
        return {'file': f_file_to_process, 'processed': False, 'err_details': e}

def _multi_join_polars(lf_dfs: List[pl.LazyFrame], on_field:str | List[str], how_m: JoinHow = JoinHow.INNER, suffix: str = "_right"):
    return reduce(lambda a,b: a.join(b, on=on_field, how=how_m, suffix=suffix), lf_dfs)

def process_wrapper(f_file_tmp: Path, only_main_leagues:bool=False):
    return asyncio.run(main_wfl(f_file_tmp, only_main_leagues))

def run_main_extract():
    f_files = get_f_list(Path('./data_countries_e'))
    to_process = []
    #f_files = f_files[:1]
    print(f"processing {f_files}")
    res_d = list()
    with ThreadPoolExecutor(os.cpu_count() // 2) as executor:
        futures = {executor.submit(process_wrapper, f_file, True): f_file for f_file in f_files if f_file.name in to_process or (not to_process)}
        for future in as_completed(futures):
            file_path = futures[future]
            try:
                result = future.result()
                if result.get('processed'):
                    res_d.append(f"[{os.getpid()}] Workflow for {os.path.basename(file_path)} completed: {result}")
                else:
                    res_d.append(
                        f"[{os.getpid()}] Workflow for {os.path.basename(file_path)} generated error : {result.get('err_details')}")
            except Exception as exc:
                res_d.append(f"[{os.getpid()}] Workflow for {os.path.basename(file_path)} generated an error when getting: {exc}")
    #
    pprint.pprint(res_d)

def combine_s_data():
    f_files = get_f_list(Path('./countries_data_processed'))
    wb_parent = load_workbook('./countries_data_processed/_data_BULK.xlsx')
    for file in f_files:
        if file.name == '_data_BULK.xlsx':
            continue
        print(f"processing file {file.name}")
        try:
            source_wb = load_workbook(file)
        except FileNotFoundError as e:
            print(f"Error: One of the files was not found. Please check the file names.")
            exit()

        source_sheet = source_wb['data']
        dest_sheet = wb_parent.create_sheet(title=re.search(r'country_processed_(.*?).xlsx', file.name).group(1).upper())

        for row in source_sheet.iter_rows():
            for cell in row:
                new_cell = dest_sheet.cell(row=cell.row, column=cell.column, value=cell.value)

                if cell.has_style:
                    new_cell.font = cell.font.copy()
                    new_cell.border = cell.border.copy()
                    new_cell.fill = cell.fill.copy()
                    new_cell.number_format = cell.number_format
                    new_cell.protection = cell.protection.copy()
                    new_cell.alignment = cell.alignment.copy()

        for col_letter, col_dim in source_sheet.column_dimensions.items():
            dest_sheet.column_dimensions[col_letter].width = col_dim.width
        source_wb.close()
    wb_parent.save(os.path.join('countries_data_processed', '_data_BULK.xlsx'))

async def _set_countries_data(client_i: httpx.AsyncClient):
    countries_link = f"https://www.transfermarkt.com/quickselect/countries"
    countries_data_r = await _simple_fetch(client_i, countries_link)
    countries_data.update({item.get("id"): item.get("name") for item in countries_data_r})

if __name__=='__main__':
    run_main_extract()