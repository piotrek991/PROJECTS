import polars as pl
import asyncio
import httpx
from lxml import etree
import requests
import json
import re
from datetime import datetime
import os
from pathlib import Path
from concurrent.futures import ProcessPoolExecutor, as_completed
import pprint
import openpyxl
from openpyxl.styles import PatternFill
import xlsxwriter
from openpyxl import load_workbook


def get_f_list(d_path: Path) -> list | None:
    if d_path.is_dir():
        f_list = list(d_path.glob('*.xlsx'))
        return f_list
    else:
        return None

def process_l_data():
    headers = {
        'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/138.0.0.0 Safari/537.36'}
    req = requests.get(
        'https://www.transfermarkt.com/laliga/startseite/wettbewerb/ES1',
        headers=headers)
    req.raise_for_status()

    h_data = etree.HTML(req.text)
    current_ses = h_data.xpath('//div[preceding-sibling::select[@name="saison_id"]]/a/span/text()')

    col_l = list()
    for column in h_data.xpath('//div[@class="responsive-table"]//thead//th'):
        if column.xpath('./a/text()'):
            col_l.append(column.xpath('./a/text()')[0])
        else:
            col_l.append(column.xpath('./text()')[0])

    print(f"produced col list : {col_l}")

    data_o_l, data_o_lt = list(), list()
    for item_l, item_lt in list(zip(h_data.xpath('//div[@class="responsive-table"]//tbody/tr'), h_data.xpath('//table[not(tfoot) and @class="items"]/tbody/tr'))):
        team_id_set = False

        data_td_t_1 = list()
        team_id = str()

        for team in item_l.xpath('./td'):
            if team.xpath('./text()') or team.xpath('./a/text()'):
                if team.xpath('./a/text()'):
                    if not team_id_set:
                        team_id = re.search(r"verein/([0-9]+)/", team.xpath('./a/@href')[0].strip()).group(1)
                        team_id_set = True
                    data_td_t_1.append(team.xpath('./a/text()')[0])
                else:
                    data_td_t_1.append(team.xpath('./text()')[0].strip())
        data_td_t_1.insert(0, team_id)

        team_id_set = False
        data_td_t_2 = list()
        team_id = str()

        for team in item_lt.xpath('./td'):
            if team.xpath('./text()') or team.xpath('./a/text()'):
                if team.xpath('./a/text()'):
                    if not team_id_set:
                        team_id = re.search(r"verein/([0-9]+)/", team.xpath('./a/@href')[0].strip()).group(1)
                        team_id_set = True
                    data_td_t_2.append(team.xpath('./a/text()')[0])
                else:
                    data_td_t_2.append(team.xpath('./text()')[0].strip())
        data_td_t_2.insert(0, team_id)

        data_o_l.append(data_td_t_1)
        data_o_lt.append(data_td_t_2)

def check_rel_list(data_l : list, ind: int):
    try:
        data_l_ind = data_l[ind].strip()
        return data_l_ind if data_l_ind else None
    except IndexError as ind_e:
        return None

async def get_process_t_data(client_i: httpx.AsyncClient, data_l: dict, async_sem_inner: asyncio.Semaphore) -> dict:
    try:
        async with async_sem_inner:
            sub_navi_base = f"https://www.transfermarkt.com/navigation/getSubNavigation?controller=verein&season=2025&id={data_l.get('team_id')}"
            sub_n_data = await client_i.get(sub_navi_base)
            sub_n_data.raise_for_status()

            sub_link = sub_n_data.json().get('results')
            sub_d_link = str()
            for sub_d_inner in sub_link:
                if sub_d_inner.get('track') == 'information_facts':
                    sub_d_link = sub_d_inner.get('open')[0].get('group')[1].get('link')

            url_base = f"https://www.transfermarkt.com{sub_d_link}"
            print(f"processing {url_base}")

            data_r = await client_i.get(url_base)
            data_r.raise_for_status()

            data_r_processed = await asyncio.to_thread(process_t_data, etree.HTML(data_r.text))
            data_r_processed.update({'team_id': data_l.get('team_id')})
            return {'process_result': True, 'fetched_data': data_r_processed}
    except httpx.HTTPStatusError as e:
        print(f"HTTP Status Error for {data_l.get('team_id')}: {e.response.status_code} - {e.response.text.strip()[:100]}...")
        return {'process_result': False, 'fetched_data': e}
    except httpx.RequestError as e:
        print(f"Request Error for {data_l.get('team_id')}: {e}")
        return {'process_result': False, 'fetched_data': e}
    except Exception as e:
        print(f"An unexpected error occurred for {data_l.get('team_id')}: {e}")
        return {'process_result': False, 'fetched_data': e}

def process_t_data(h_data) ->  dict:
    data_o = dict()
    props_d = dict()

    try:
        league_level = h_data.xpath('//div[@class="data-header__club-info"]/span[@class="data-header__label"][1]/span/a/text()')[1].strip()
    except IndexError as ind_e:
        league_level = 'NOT DEFINED'

    for item in h_data.xpath('//table[preceding-sibling::h2[text()="All titles"]]/tbody/tr'):
        data_td_det = list()
        for team in item.xpath('./td'):
            if check_rel_list(team.xpath('./text()'), 0) or check_rel_list(team.xpath('./a/text()'),0):
                if team.xpath('./a/text()'):
                    data_td_det.append(team.xpath('./a/text()')[0])
                else:
                    data_td_det.append(team.xpath('./text()')[0].strip())
        if len(data_td_det) == 2:
            if not (item:=data_o.get(data_td_det[1])):
                data_o.update({data_td_det[1]: [data_td_det[0]]})
            else:
                item.append(data_td_det[0])
                data_o.update({data_td_det[1]: item})

    for item in h_data.xpath('//div[@class="data-header__info-box "]/*/ul'):
        for item_i in item.xpath('./li'):
            li_data_prop = item_i.xpath('./text()')[0].strip()
            if check_rel_list( item_i.xpath('./span/text()'), 0):
                li_data_d = item_i.xpath('./span/text()')[0].strip()
            elif check_rel_list(item_i.xpath('./span/a/text()'), 0) and not check_rel_list(item_i.xpath('./span/a[following-sibling::span]/text()'), 0):
                li_data_d =  item_i.xpath('./span/a/text()')[0].strip()
            else:
                prop_d = str()
                if check_rel_list(item_i.xpath('./span/a/text()') ,0):
                    prop_d = prop_d + ', ' + item_i.xpath('./span/a/text()')[0].strip() if prop_d else item_i.xpath('./span/a/text()')[0].strip()

                if check_rel_list(item_i.xpath('./span/span/text()'), 0):
                    prop_d = prop_d + ', ' + item_i.xpath('./span/span/text()')[0].strip() if prop_d else  item_i.xpath('./span/span/text()')[0].strip()
                else:
                    prop_d = prop_d + ', ' + item_i.xpath('./span/span/a/text()')[0].strip() if prop_d else \
                        item_i.xpath('./span/span/a/text()')[0].strip()
                li_data_d = prop_d
            props_d.update({li_data_prop: li_data_d})

    multipliers = {
        'k': 10 ** 3
        , 'm': 10 ** 6
        , 'bn': 10 ** 9
    }
    try:
        currency_multiplier = h_data.xpath('//div[@class="data-header__box--small"]//span/text()')[1].strip()
        total_team_market = h_data.xpath('//div[@class="data-header__box--small"]/a/text()')[0].strip()
    except IndexError as ind_e:
        currency_multiplier = 'NOT DEFINED'
        total_team_market = '0'
    return {**props_d, 'promotions_history': json.dumps(data_o), 'league_level': league_level, 'total_m_value': f"{float(total_team_market) * multipliers.get(currency_multiplier, 1)}"}

def map_year(y_two_digit: str) -> int:
    d_now = datetime.now().year
    two_digit = d_now % 100

    if int(y_two_digit) > two_digit:
        return 1900 + int(y_two_digit)
    else:
        return 2000 + int(y_two_digit)

async def main_wfl(f_file_to_process: Path):
    async_sem = asyncio.Semaphore(5)
    try:
        country_nm = re.search(r'data_(.*?).xlsx', str(f_file_to_process.name)).group(1)
        print(f"processing country_nm {country_nm}")
        pl_df = pl.read_excel(f_file_to_process, sheet_name=country_nm).select('team_id', 'team_name', 'team_link').unique()
        pl_df_r = pl_df.select(pl.struct(['team_id', 'team_link']).alias('t_data')).to_series().to_list()
        page_headers = {'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/138.0.0.0 Safari/537.36'}

        async with httpx.AsyncClient(headers=page_headers, timeout=60) as client:
            tasks = list()
            for team_t in pl_df_r:
                tasks.append(asyncio.create_task(get_process_t_data(client, team_t, async_sem)))
            data_f = await asyncio.gather(*tasks, return_exceptions=True)

        data_f_processed = list()
        for item in data_f:
            if item.get('process_result'):
                data_f_processed.append(item.get('fetched_data'))

        pl_df_processed = pl.DataFrame(data_f_processed)
        pl_df_processed = pl_df_processed.with_columns(
            pl.col('promotions_history').map_elements(
                lambda row: json.loads(row) if row else None
                , return_dtype=pl.Object
            ),
        )

        pl_df_comb = pl_df.join(pl_df_processed, on='team_id')
        pl_df_comb = pl_df_comb.with_columns(
            pl.col('promotions_history').map_elements(
                lambda d: list(d.items()) if isinstance(d, dict) else None,
                return_dtype=pl.List(pl.Struct([pl.Field('ach_name', pl.Utf8), pl.Field('ach_years', pl.List(pl.Utf8))]))
            ).alias('promotions_history_k'),
        ).explode('promotions_history_k').unnest('promotions_history_k').explode('ach_years') \
            .with_columns(
            pl.col('ach_years').map_elements(
                lambda row: map_year(re.search(r'^([0-9]+)', row).group(1)), return_dtype=pl.Int16
            ).alias('ach_year_int')
        ).sort(["team_id", "ach_year_int"], descending=True) \
            .with_columns(
            (pl.col('ach_years') + "-" + pl.col('ach_name')).alias('Promotion History')
        ).select(
            pl.exclude("team_id", "team_link", "Current transfer record:", "promotions_history", "ach_name", "ach_years",
                       "ach_year_int")) \
            .rename({'team_name': 'Team Name', 'Squad size:': 'Squad Size', 'Average age:': 'Average Age',
                     'Foreigners:': 'Foreigners', 'Stadium:': 'Stadium', 'league_level': 'League Level', 'National team players:': 'National Team Players' ,'total_m_value': 'Total Market Value'})

        pl_cols_group = list(set(pl_df_comb.columns).difference({'Promotion History'}))
        pl_df_comb_group = pl_df_comb.group_by(pl_cols_group).agg(
            pl.col('Promotion History').str.join(",\n").alias("Promotion History")
        ).with_columns(
            pl.col("Foreigners").str.split(",").map_elements(lambda row: row.str.strip_chars(), return_dtype=pl.List(pl.String)).list.to_struct(fields=['F num', 'F %'])
        ).unnest("Foreigners").with_columns(
            pl.when(pl.col('F num').str.replace('-','0') == pl.lit('0')).then(pl.lit('0')).otherwise(pl.col('F %').str.replace_all(r'[% ]', '')).cast(pl.Float32).alias('Foreigners %'),
            pl.col('F num').str.replace('-', '0').cast(pl.Int64).alias('Foreigners Num'),
            pl.col('Average Age').cast(pl.Float32).round(2),
            pl.col('National Team Players').cast(pl.Int64),
            pl.col('Squad Size').cast(pl.Int64),
            pl.col('Total Market Value').cast(pl.Float32),
        ).lazy()
        pl_df_comb_group = pl_df_comb_group.select('Team Name', 'League Level', 'Total Market Value', 'Squad Size', 'Average Age', 'National Team Players','Foreigners Num', 'Foreigners %', 'Stadium', 'Promotion History').collect()
        f_name_base = f"country_processed_{country_nm}.xlsx"

        os.makedirs('countries_data_processed', exist_ok=True)
        pl_df_comb_group.write_excel(
            os.path.join('countries_data_processed', f_name_base)
            , 'data'
            ,  dtype_formats={
                pl.Float32: "#,##0.00"
            }
        )

        wb = openpyxl.load_workbook(os.path.join('countries_data_processed', f_name_base))
        sheet = wb.active
        pattern = PatternFill(start_color='ADD8E6', end_color='ADD8E6', fill_type='solid')
        for cell in sheet[1]:
            cell.fill = pattern

        sheet.freeze_panes = 'A2'
        for loc, letter in enumerate(['A','B','C','D','E','F','G','H','I','J'], start=1):
            m_width = 0
            for r_num in range(1, sheet.max_row + 1):
                if letter == 'C':
                    sheet.cell(row=r_num, column=loc).number_format = '#,##0.00" €"'

                if sheet[f'{letter}{r_num}'].value:
                    if len(str(sheet[f'{letter}{r_num}'].value)) > m_width:
                        m_width = len(str(sheet[f'{letter}{r_num}'].value))
            sheet.column_dimensions[letter].width = m_width
        wb.save(os.path.join('countries_data_processed', f_name_base))
        wb.close()

        return { 'file': f_file_to_process, 'processed': True}

    except Exception as e:
        return { 'file': f_file_to_process, 'processed': False, 'err_details': e }

def process_wrapper(f_file_tmp: Path):
    return asyncio.run(main_wfl(f_file_tmp))

def run_main_extract():
    f_files = get_f_list(Path('./data_countries_e'))
    res_d = list()
    with ProcessPoolExecutor(max_workers=3) as executor:
        futures = {executor.submit(process_wrapper, f_file): f_file for f_file in f_files}
        for future in as_completed(futures):
            file_path = futures[future]
            try:
                result = future.result()
                if result.get('processed'):
                    res_d.append(f"[{os.getpid()}] Workflow for {os.path.basename(file_path)} completed: {result}")
                else:
                    res_d.append(
                        f"[{os.getpid()}] Workflow for {os.path.basename(file_path)} generated error : {result.get('err_details')}")
            except Exception as exc:
                res_d.append(f"[{os.getpid()}] Workflow for {os.path.basename(file_path)} generated an error: {exc}")

    pprint.pprint(res_d)

def combine_s_data():
    f_files = get_f_list(Path('./countries_data_processed'))
    wb_parent = load_workbook('./countries_data_processed/_data_BULK.xlsx')
    for file in f_files:
        if file.name == '_data_BULK.xlsx':
            continue
        print(f"processing file {file.name}")
        try:
            source_wb = load_workbook(file)
        except FileNotFoundError as e:
            print(f"Error: One of the files was not found. Please check the file names.")
            exit()

        source_sheet = source_wb['data']
        dest_sheet = wb_parent.create_sheet(title=re.search(r'country_processed_(.*?).xlsx', file.name).group(1).upper())

        for row in source_sheet.iter_rows():
            for cell in row:
                new_cell = dest_sheet.cell(row=cell.row, column=cell.column, value=cell.value)

                if cell.has_style:
                    new_cell.font = cell.font.copy()
                    new_cell.border = cell.border.copy()
                    new_cell.fill = cell.fill.copy()
                    new_cell.number_format = cell.number_format
                    new_cell.protection = cell.protection.copy()
                    new_cell.alignment = cell.alignment.copy()

        for col_letter, col_dim in source_sheet.column_dimensions.items():
            dest_sheet.column_dimensions[col_letter].width = col_dim.width
        source_wb.close()
    wb_parent.save(os.path.join('countries_data_processed', '_data_BULK.xlsx'))
if __name__=='__main__':
   combine_s_data()